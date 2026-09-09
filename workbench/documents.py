import io
import zipfile
from pathlib import Path

MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = 300_000


def parse_document(name: str, content: bytes) -> str:
    if len(content) > MAX_BYTES:
        raise ValueError("文件不能超过 20MB。")
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".txt":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("gb18030")
        elif suffix == ".docx":
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("文档解压后过大。")
            document = Document(io.BytesIO(content))
            text = "\n".join([p.text for p in document.paragraphs] + [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows])
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("PDF 已加密，请先解除密码保护。")
            parts, total = [], 0
            for page in reader.pages:
                part = page.extract_text() or ""
                total += len(part)
                if total > MAX_CHARS:
                    raise ValueError("提取文本不能超过 30 万字符。")
                parts.append(part)
            text = "\n".join(parts)
        else:
            raise ValueError("仅支持 TXT、DOCX 和 PDF 文件。")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("文件解析失败，请检查文件是否损坏或编码是否正确。") from exc
    text = text.strip()
    if not text:
        raise ValueError("未提取到文字；扫描 PDF 请先进行 OCR 或改用文字文档。")
    if len(text) > MAX_CHARS:
        raise ValueError("提取文本不能超过 30 万字符。")
    return text
