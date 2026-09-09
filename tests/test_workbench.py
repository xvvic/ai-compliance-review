import asyncio
import io
import json
import os
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from server import create_app
from workbench.config import ConfigStore, Settings, agent_environment
from workbench.documents import MAX_BYTES, MAX_CHARS, parse_document
from workbench.jobs import ReviewManager, scan


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path)) as client:
        client.headers["x-session-token"] = client.get('/api/session').json()['token']
        yield client


def test_boundary_and_no_secret_echo(client):
    settings = {'secret': 'private-test-value', 'model': 'sonnet'}
    assert client.put('/api/config', json=settings, headers={'origin': 'https://outside.example'}).status_code == 403
    assert client.put('/api/config', json=settings, headers={'x-session-token': ''}).status_code == 403
    response = client.put('/api/config', json=settings)
    assert response.status_code == 200
    assert 'private-test-value' not in response.text
    assert 'private-test-value' not in client.app.state.config.path.read_text()
    assert client.app.state.config.load().secret == settings['secret']
    assert 'private-test-value' not in client.get('/api/config').text
    assert client.get('/api/config', headers={'host': 'outside.example'}).status_code == 403


def test_configuration_preserve_clear_and_validation(client):
    assert client.put('/api/config', json={'secret': 'sample-secret'}).status_code == 200
    assert client.put('/api/config', json={'model': 'another-model'}).json()['has_secret']
    assert not client.put('/api/config', json={'clear_secret': True}).json()['has_secret']
    bad = client.put('/api/config', json={'base_url': 'https://someone:password@example.com', 'secret': 'dont-echo-this'})
    assert bad.status_code == 422
    assert 'dont-echo-this' not in bad.text and 'password' not in bad.text


@pytest.mark.parametrize('name,content', [('empty.txt', b''), ('bad.docx', b'bad'), ('bad.pdf', b'bad'), ('x.exe', b'not supported'), ('huge.txt', b'a' * (MAX_BYTES + 1)), ('long.txt', b'a' * (MAX_CHARS + 1))], ids=['empty', 'docx', 'pdf', 'extension', 'bytes-limit', 'text-limit'])
def test_invalid_documents(name, content):
    with pytest.raises(ValueError):
        parse_document(name, content)


def test_text_encodings_and_docx_tables():
    assert parse_document('a.txt', '中文材料'.encode('utf-8-sig')) == '中文材料'
    assert parse_document('a.txt', '中文材料'.encode('gb18030')) == '中文材料'
    doc = Document()
    doc.add_paragraph('企业方案')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格里的个人信息'
    buffer = io.BytesIO()
    doc.save(buffer)
    assert '表格里的个人信息' in parse_document('a.docx', buffer.getvalue())


def test_scanned_and_encrypted_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match='OCR'):
        parse_document('scan.pdf', buffer.getvalue())
    writer.encrypt('test-password')
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match='加密'):
        parse_document('locked.pdf', buffer.getvalue())


def test_upload_missing_config_and_example(client):
    result = client.post('/api/documents/parse', files={'file': ('material.txt', '企业方案'.encode(), 'text/plain')})
    assert result.status_code == 200 and result.json()['chars'] == 4
    assert client.post('/review/stream', json={'document_text': 'test'}).status_code == 400
    assert client.get('/api/example').json()['example'] is True
    assert client.app.state.manager.current is None


def completed_job():
    return {'id': 'abc123', 'status': 'completed', 'reviewed': False, 'events': [], 'report': {'generated_at': '2026-01-01', 'risk_scan': {'matched_rules': [{'id': 'R1', 'final_level': 'L3', 'category': 'data', 'risk_type': 'risk'}]}}}


def test_decisions_linkage_and_idempotence(client, tmp_path):
    client.app.state.manager.current = completed_job()
    item = {'rule_id': 'R1', 'action': '调整等级', 'review_level': 'L2', 'evidence_note': ''}
    assert client.post('/api/review/abc123/decisions', json={'items': [item]}).status_code == 400
    item['evidence_note'] = 'Evidence checked'
    assert client.post('/api/review/other/decisions', json={'items': [item]}).status_code == 409
    assert client.post('/api/review/abc123/decisions', json={'items': [item, item]}).status_code == 400
    assert client.post('/api/review/abc123/decisions', json={'items': [item]}).status_code == 200
    row = json.loads((tmp_path / 'reviews/review_abc123.jsonl').read_text())
    assert row['adjusted'] is True and row['review_id'] == 'abc123' and row['system_level'] == 'L3'
    assert client.post('/api/review/abc123/decisions', json={'items': [item]}).status_code == 409


def test_decision_write_failure_does_not_mark_reviewed(client, monkeypatch):
    client.app.state.manager.current = completed_job()
    monkeypatch.setattr('server.atomic_write', lambda *args: (_ for _ in ()).throw(OSError('private-path')))
    response = client.post('/api/review/abc123/decisions', json={'items': [{'rule_id': 'R1', 'action': '认可初评', 'review_level': 'L3'}]})
    assert response.status_code == 500
    assert not client.app.state.manager.current['reviewed']
    assert 'private-path' not in response.text


class FakeInput:
    def write(self, value): pass
    async def drain(self): pass
    def close(self): pass


class FakeProcess:
    def __init__(self, events):
        self.stdin = FakeInput()
        self.stdout = asyncio.StreamReader()
        for event in events:
            self.stdout.feed_data((json.dumps(event) + '\n').encode())
        self.stdout.feed_eof()
        self.returncode = 0
    async def wait(self): return 0


@pytest.mark.asyncio
@pytest.mark.parametrize('events,status', [([{'type': 'error', 'content': 'model failed'}, {'type': 'final', 'content': 'must not publish'}], 'failed'), ([{'type': 'final', 'content': '# report'}], 'completed'), ([], 'failed')])
async def test_stream_terminal_states(tmp_path, monkeypatch, events, status):
    async def spawn(*args, **kwargs): return FakeProcess(events)
    monkeypatch.setattr(asyncio, 'create_subprocess_exec', spawn)
    manager = ReviewManager(tmp_path)
    await manager.start('个人信息跨境', 'test.txt', Settings(secret='test-key'))
    await manager.task
    assert manager.current['status'] == status
    types = [event['type'] for event in manager.current['events']]
    assert types[-1] == 'done'
    if status == 'failed':
        assert 'final' not in types and not list(tmp_path.glob('reports/*'))
    else:
        assert types.count('final') == 1 and manager.current['saved']


@pytest.mark.asyncio
async def test_report_save_failure(tmp_path, monkeypatch):
    async def spawn(*args, **kwargs): return FakeProcess([{'type': 'final', 'content': '# report'}])
    monkeypatch.setattr(asyncio, 'create_subprocess_exec', spawn)
    monkeypatch.setattr('workbench.jobs.atomic_write', lambda *args: (_ for _ in ()).throw(OSError('secret path')))
    manager = ReviewManager(tmp_path)
    await manager.start('材料', 'a.txt', Settings())
    await manager.task
    assert manager.current['status'] == 'completed' and not manager.current['saved']
    assert manager.current['save_error'] and manager.current['report']


@pytest.mark.asyncio
async def test_single_job_and_cancel(tmp_path, monkeypatch):
    async def spawn(*args, **kwargs):
        process = FakeProcess([])
        process.stdout = asyncio.StreamReader()
        return process
    monkeypatch.setattr(asyncio, 'create_subprocess_exec', spawn)
    manager = ReviewManager(tmp_path)
    await manager.start('材料', 'a.txt', Settings())
    await asyncio.sleep(0.1)
    with pytest.raises(ValueError):
        await manager.start('材料', 'b.txt', Settings())
    await manager.cancel()
    assert manager.current['status'] == 'cancelled'
    assert manager.current['events'][-1]['type'] == 'done'


def test_scan_degradation_and_environment(monkeypatch, tmp_path):
    assert scan('个人信息跨境')['matched_rules']
    monkeypatch.setattr('workbench.jobs.PLUGIN', tmp_path / 'missing')
    assert scan('个人信息')['error']
    env = agent_environment(Settings(secret='test-key'), tmp_path)
    assert env['ANTHROPIC_AUTH_TOKEN'] == ''
    assert env['CLAUDE_CONFIG_DIR'].startswith(str(tmp_path))


def test_config_recovery_and_endpoint_change(tmp_path):
    config = ConfigStore(tmp_path)
    config.path.write_text('broken json')
    assert config.public()['config_error']
    config.save(Settings(secret='first-key'))
    assert config.load().secret == 'first-key'
    config.save(Settings(base_url='https://other.example'))
    assert not config.load().secret


def test_connection_uses_agent_and_sanitizes_failures(client, monkeypatch):
    seen = []
    async def success(settings):
        seen.append(settings.model)
    monkeypatch.setattr('workbench.connection.test_connection', success)
    assert client.post('/api/config/test', json={'secret': 'fake-key', 'model': 'sonnet'}).status_code == 200
    assert seen == ['sonnet']
    async def failure(settings):
        raise RuntimeError('private-secret private-path')
    monkeypatch.setattr('workbench.connection.test_connection', failure)
    response = client.post('/api/config/test', json={'secret': 'fake-key'})
    assert response.status_code == 400 and 'private-secret' not in response.text


@pytest.mark.asyncio
async def test_cancel_before_task_runs(tmp_path):
    manager = ReviewManager(tmp_path)
    await manager.start('test', 'a.txt', Settings())
    await manager.cancel()
    assert manager.current['status'] == 'cancelled'
    assert manager.current['events'][-1]['type'] == 'done'
